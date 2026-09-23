"""
Activa IT - Descargador automático de Cartas Glosas / Liquidaciones SOAT
============================================================================
Un solo programa, un solo puerto, DOS páginas independientes:
  - /estado  -> SIS / Seguros del Estado  (soatestado.sis.co)
  - /sura    -> Suramericana              (soatsura.sis.co)

Cada página tiene su propio panel, su propio job en segundo plano, su propia
carpeta de descargas (downloads/<empresa>/<IPS>/...) y su propio progreso —
se pueden correr las dos AL MISMO TIEMPO sin que se pisen.

La lógica (login, navegación, llenado del formulario, captura del PDF,
reintento automático) es exactamente la misma para las dos; lo único que
cambia por empresa es la URL de login, el mapa de IPS/NIT, el mapa de ciudad
de login, y el logo/colores — todo eso vive en EMPRESAS al inicio del archivo.
"""

import os
import re
import json
import csv
import time
import base64
import zipfile
import unicodedata
import threading
import logging
from io import BytesIO
from datetime import datetime, timezone
from pathlib import Path
from flask import Blueprint, render_template, request, jsonify, send_file, send_from_directory, abort
from . import concurrency
from . import historial_db
from . import registro_rutas
from .catalogo_ips import resolver, validar_identidad, normalizar_nit, listar_catalogo_por_responsable

try:
    import openpyxl
    from openpyxl.styles import Font
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

bp = Blueprint("estado_sura", __name__)

BASE_DIR = Path(__file__).resolve().parent.parent  # raíz del proyecto (un nivel arriba de bots/)
PORT = int(os.environ.get("PORT", 8080))
DOWNLOAD_DIR = BASE_DIR / "downloads"
DOWNLOAD_DIR.mkdir(exist_ok=True)


# ==================== CONFIG POR EMPRESA ====================
# Todo lo que distingue a un portal del otro vive aquí. Para agregar una
# tercera aseguradora en el futuro, solo hay que sumar otra entrada a este
# diccionario (y su logo en static/) — el resto del código no cambia.
EMPRESAS = {
    "estado": {
        "nombre": "SIS Estado",
        "titulo": "Descargador Cartas Glosas SIS VIDA - ESTADO",
        "subtitulo": "Activa IT · Automatización de descarga de soportes (Seguros del Estado)",
        "login_url": "https://soatestado.sis.co/#/auth/login",
        "dashboard_dominio": "https://soatestado.sis.co/",
        "logo": "logo_estado.png",
        "zip_prefix": "liquidaciones",
        "accent": "#a41e35",
        "accent_hover": "#841829",
        "mapa_ips": {
            "900267064": "CLINICA_BAHIA",
            "900513306": "FUNDACION_MARIA_REINA_SUCRE_SINCELEJO",
            "900600550": "CLINICA_BARU",
            "900631361": "VALLE_SALUD_NORTE",          # o VALLE_SALUD_SUR (mismo NIT)
            "900657731": "CMR_BAHIA",
            "900827065": "CDI_BAHIA",
            "900847382": "CMR_VALLESALUD",
            "900954800": "CMR_BARU",
            "900257333": "ODONTOTRANS",
            "900792417": "RUC_PACIFICA",
            "900826509": "RUC_MAGDALENA",
            "900900754": "SAN_FERNANDO",
            "901081281": "URGETRAUMA_SAN_FERNANDO",    # o ALVERNIA_TULUA (mismo NIT)
            "800255591": "MEDICO_QUIRURGICA_TULUA",
            "900002780": "FUNDACION_CAMPBELL",
            "900558595": "FUNDACION_MEDICA_CAMPBELL",
            "901149757": "UNIDAD_MEDICA_DE_TRAUMA_VALLE_SALUD",
            "900469882": "CENTRO_MEDICO_SERVISALUD_INTEGRAL_IPS",
            "901523868": "MOVID_IPS",
            "802024329": "RED_DE_URGENCIA_DE_LA_COSTA",
        },
        "ciudad_login": {
            "900267064": "SANTA MARTA - MAGDALENA",
            "900657731": "SANTA MARTA - MAGDALENA",
            "900827065": "SANTA MARTA - MAGDALENA",
            "900826509": "SANTA MARTA - MAGDALENA",
            "900513306": "SINCELEJO - SUCRE",
            "900600550": "CARTAGENA DE INDIAS - BOLIVAR",
            "900954800": "CARTAGENA DE INDIAS - BOLIVAR",
            "900631361": "CALI - VALLE DEL CAUCA",
            "900847382": "CALI - VALLE DEL CAUCA",
            "900257333": "CALI - VALLE DEL CAUCA",
            "900792417": "CALI - VALLE DEL CAUCA",
            "900900754": "CALI - VALLE DEL CAUCA",
            "901081281": "CALI - VALLE DEL CAUCA",
            "800255591": "CALI - VALLE DEL CAUCA",
            "901149757": "CALI - VALLE DEL CAUCA",
            "900469882": "JAMUNDI - VALLE DEL CAUCA",
            "900002780": "BARRANQUILLA - ATLANTICO",
            "901523868": "BARRANQUILLA - ATLANTICO",
            "900558595": "MALAMBO - ATLANTICO",
            "802024329": "BARRANQUILLA - ATLANTICO",
        },
        # NITs con más de una sede: el usuario debe escoger la ciudad en el panel.
        "sedes_login": {
            "900002780": ["BARRANQUILLA - ATLANTICO", "MALAMBO - ATLANTICO",
                          "SOLEDAD - ATLANTICO", "BARANOA - ATLANTICO",
                          "SABANALARGA - ATLANTICO"],
            "900558595": ["MALAMBO - ATLANTICO", "SOLEDAD - ATLANTICO"],
        },
    },
    "sura": {
        "nombre": "Suramericana",
        "titulo": "Descargador Soportes SURAMERICANA",
        "subtitulo": "Activa IT · Automatización de descarga de soportes",
        "login_url": "https://soatsura.sis.co/#/auth/login",
        "dashboard_dominio": "https://soatsura.sis.co/",
        "logo": "logo_sura.png",
        "zip_prefix": "soportes",
        "accent": "#0032a1",
        "accent_hover": "#00227a",
        # Mismo mapa que Estado por ahora (edítalo si Suramericana usa
        # NITs/IPS distintos — el resto del código no cambia).
        "mapa_ips": {
            "900267064": "CLINICA_BAHIA",
            "900513306": "FUNDACION_MARIA_REINA_SUCRE_SINCELEJO",
            "900600550": "CLINICA_BARU",
            "900631361": "VALLE_SALUD_NORTE",
            "900657731": "CMR_BAHIA",
            "900827065": "CDI_BAHIA",
            "900847382": "CMR_VALLESALUD",
            "900954800": "CMR_BARU",
            "900257333": "ODONTOTRANS",
            "900792417": "RUC_PACIFICA",
            "900826509": "RUC_MAGDALENA",
            "900900754": "SAN_FERNANDO",
            "901081281": "URGETRAUMA_SAN_FERNANDO",
            "800255591": "MEDICO_QUIRURGICA_TULUA",
            "900002780": "FUNDACION_CAMPBELL",
            "900558595": "FUNDACION_MEDICA_CAMPBELL",
            "901149757": "UNIDAD_MEDICA_DE_TRAUMA_VALLE_SALUD",
            "900469882": "CENTRO_MEDICO_SERVISALUD_INTEGRAL_IPS",
            "901523868": "MOVID_IPS",
            "802024329": "RED_DE_URGENCIA_DE_LA_COSTA",
        },
        "ciudad_login": {
            "900267064": "SANTA MARTA - MAGDALENA",
            "900657731": "SANTA MARTA - MAGDALENA",
            "900827065": "SANTA MARTA - MAGDALENA",
            "900826509": "SANTA MARTA - MAGDALENA",
            "900513306": "SINCELEJO - SUCRE",
            "900600550": "CARTAGENA DE INDIAS - BOLIVAR",
            "900954800": "CARTAGENA DE INDIAS - BOLIVAR",
            "900631361": "CALI - VALLE DEL CAUCA",
            "900847382": "CALI - VALLE DEL CAUCA",
            "900257333": "CALI - VALLE DEL CAUCA",
            "900792417": "CALI - VALLE DEL CAUCA",
            "900900754": "CALI - VALLE DEL CAUCA",
            "901081281": "CALI - VALLE DEL CAUCA",
            "800255591": "CALI - VALLE DEL CAUCA",
            "901149757": "CALI - VALLE DEL CAUCA",
            "900469882": "JAMUNDI - VALLE DEL CAUCA",
            "900002780": "BARRANQUILLA - ATLANTICO",
            "901523868": "BARRANQUILLA - ATLANTICO",
            "900558595": "MALAMBO - ATLANTICO",
            "802024329": "BARRANQUILLA - ATLANTICO",
        },
        # NITs con más de una sede: el usuario debe escoger la ciudad en el panel.
        "sedes_login": {
            "900002780": ["BARRANQUILLA - ATLANTICO", "MALAMBO - ATLANTICO",
                          "SOLEDAD - ATLANTICO", "BARANOA - ATLANTICO",
                          "SABANALARGA - ATLANTICO"],
            "900558595": ["MALAMBO - ATLANTICO", "SOLEDAD - ATLANTICO"],
        },
    },
}


def empresa_o_404(empresa):
    if empresa not in EMPRESAS:
        abort(404, description=f"Empresa desconocida: {empresa}")
    return EMPRESAS[empresa]


# ==================== ESTADO GLOBAL (multi-cuenta: un job por empresa+NIT) ====================
def _estado_inicial():
    return {
        "running": False,
        "stopping": False,
        "logs": [],
        "stats": {"total": 0, "descargadas": 0, "errores": 0},
        "finished": False,
        "error": None,
        "errores_detalle": [],
        "descargas_exitosas": [],
        "errores_excel_url": None,
        "zip_url": None,
    }


MAX_CONCURRENT_JOBS = int(os.environ.get("MAX_CONCURRENT_JOBS", "2"))
jobs = {}  # "empresa:empresa_id" -> {"state":..., "lock":..., "browser":...}
jobs_registry_lock = threading.RLock()


def resolve_empresa_id(usuario: str) -> str:
    """El identificador de la cuenta es el NIT incluido en el usuario (o el
    propio usuario en mayúsculas si no se encuentra un NIT de 9-12 dígitos)."""
    m = re.search(r"(\d{9,12})", usuario or "")
    return m.group(1) if m else (usuario or "").strip().upper()


def _job_key(empresa, empresa_id):
    return f"{empresa}:{empresa_id}"


def get_or_create_job(empresa, empresa_id):
    key = _job_key(empresa, empresa_id)
    with jobs_registry_lock:
        if key not in jobs:
            jobs[key] = {"state": _estado_inicial(), "lock": threading.Lock(), "browser": None, "ips_nombre": None}
        return jobs[key]


def get_job_or_none(empresa, empresa_id):
    with jobs_registry_lock:
        return jobs.get(_job_key(empresa, empresa_id))


def count_running_jobs(empresa):
    with jobs_registry_lock:
        return sum(1 for k, j in jobs.items() if k.startswith(f"{empresa}:") and j["state"]["running"])


# ==================== UTILIDADES ====================
def sanitizar_nombre(nombre: str) -> str:
    nombre = str(nombre).strip().upper()
    nombre = unicodedata.normalize("NFKD", nombre).encode("ascii", "ignore").decode()
    nombre = re.sub(r"[^A-Z0-9]+", "_", nombre).strip("_")
    return nombre or "IPS_DESCONOCIDA"


def sanitizar_factura(valor) -> str:
    texto = str(valor).strip()
    texto = re.sub(r'[\\/*?:"<>|]', "", texto)
    return texto or "SIN_FACTURA"


SINIESTRO_RE = re.compile(r"^\s*(\d+)\s*/\s*(\d{2,4})\s*\*\s*(\d+)\s*$")


def parse_siniestro(valor):
    """'17155/2026*3' -> {'numero': '17155', 'anio': '2026', 'cuenta': '3'}"""
    m = SINIESTRO_RE.match(str(valor))
    if not m:
        return None
    numero, anio, cuenta = m.groups()
    if len(anio) == 2:
        anio = "20" + anio
    return {"numero": numero, "anio": anio, "cuenta": cuenta}


def extraer_ips_desde_usuario(usuario: str, mapa_ips: dict):
    m = re.search(r"(\d{9,12})", usuario)
    nit = m.group(1) if m else None
    nombre = mapa_ips.get(nit)
    if nombre:
        return nombre, nit
    return sanitizar_nombre(usuario or "IPS_DESCONOCIDA"), nit


def log(job, empresa, msg, level="info"):
    ts = datetime.now().strftime("%H:%M:%S")
    if job is not None:
        with job["lock"]:
            job["state"]["logs"].append({"ts": ts, "msg": msg, "level": level})
    (logger.error if level == "error" else logger.info)(f"[{empresa}] {msg}")
    if level == "error" and job is not None:
        try:
            historial_db.registrar_error_ejecucion(
                job["state"].get("ejecucion_id"), None, empresa, EMPRESAS.get(empresa, {}).get("nombre"),
                job["state"].get("ips_identity"), "bot", msg, "log", True, 1, "error"
            )
        except Exception:
            pass


def reset_state(job):
    with job["lock"]:
        job["state"].update(_estado_inicial())


def stop_job(job, empresa):
    with job["lock"]:
        job["state"]["stopping"] = True
    log(job, empresa, "🛑 Solicitando detención del proceso...", "warn")
    browser = job.get("browser")
    if browser:
        try:
            browser.close()
        except Exception as e:
            log(job, empresa, f"  → Error al cerrar navegador: {e}", "error")


# ==================== PERSISTENCIA (REANUDACIÓN) ====================
def cargar_progreso(ips_dir: Path):
    p = ips_dir / "progreso.json"
    if p.exists():
        try:
            return set(json.load(open(p, encoding="utf-8")).get("completadas", []))
        except Exception:
            pass
    return set()


def guardar_progreso(ips_dir: Path, exitosas, meta=None):
    """
    Guarda no solo la lista de facturas completadas (como antes, para que
    cargar_progreso() siga reanudando un lote a medio camino exactamente
    igual), sino también el siniestro y la fecha real de cada una, más
    metadata (identidad, ips, aseguradora) — necesario para poder migrar
    esto a la base de datos de historial más adelante sin perder detalle.
    """
    p = ips_dir / "progreso.json"
    try:
        detalle = {
            e["factura"]: {
                "siniestro": f"{e['numero']}/{e['anio']}*{e['cuenta']}",
                "fecha_descarga": e.get("timestamp") or datetime.now(timezone.utc).isoformat(),
            }
            for e in exitosas
        }
        data = {
            "completadas": list(detalle.keys()),
            "detalle": detalle,
            "actualizado": datetime.now(timezone.utc).isoformat(),
        }
        if meta:
            data["meta"] = meta
        json.dump(data, open(p, "w", encoding="utf-8"), indent=2)
    except Exception as e:
        logger.warning(f"Error al escribir progreso.json: {e}")
        return

    try:
        historial_db.registrar_descargas(
            aseguradora=(meta or {}).get("aseguradora", "Desconocida"),
            ips_nombre=(meta or {}).get("ips_nombre", "IPS_NO_IDENTIFICADA"),
            periodo=(meta or {}).get("periodo"), identidad=(meta or {}).get("identidad"),
            items=[{"factura": e["factura"], "siniestro": f"{e['numero']}/{e['anio']}*{e['cuenta']}", "fecha_descarga": e.get("timestamp")} for e in exitosas],
            ips=(meta or {}).get("ips"), sede=(meta or {}).get("sede"),
        )
    except Exception as e:
        logger.warning(f"Error al registrar en el historial (tabla descargas): {e}")

    try:
        historial_db.registrar_facturas_ejecucion(
            (meta or {}).get("ejecucion_id"), "estado_sura", (meta or {}).get("aseguradora"),
            (meta or {}).get("ips"), (meta or {}).get("periodo"),
            [{"factura": e["factura"], "estado": "exitosa", "archivo": e.get("archivo"), "fecha_fin": e.get("timestamp")} for e in exitosas],
        )
    except Exception as e:
        logger.warning(f"Error al registrar la ejecución (tabla ejecucion_facturas): {e}")


# ==================== REPORTE Y ZIP ====================
def generar_reporte_excel(ips_dir: Path, ips_nombre: str, exitosas, errores):
    if not EXCEL_AVAILABLE:
        return None
    path = ips_dir / f"reporte_{ips_nombre}.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Descargadas"
    ws1.append(["N° Factura", "N° Siniestro", "Año", "N° Cuenta", "Archivo", "Fecha/Hora"])
    for e in exitosas:
        ws1.append([e["factura"], e["numero"], e["anio"], e["cuenta"], e["archivo"], e["timestamp"]])
    ws2 = wb.create_sheet("Errores")
    ws2.append(["N° Factura", "Siniestro", "Error", "Captura", "Fecha/Hora"])
    for e in errores:
        ws2.append([e["factura"], e["siniestro_raw"], e["error"], e.get("captura", ""), e["timestamp"]])
    wb.save(path)
    return path


def generar_excel_solo_errores(dl_dir_empresa: Path, ips_nombre: str, errores):
    """
    Excel aparte, SOLO con las facturas que quedaron en error después de
    agotar los reintentos programados. Pensado para descarga automática
    (no bloquea ni afecta el flujo principal: cualquier fallo aquí se
    ignora en silencio, ya que el reporte combinado ya tiene esta misma
    información como respaldo).
    """
    if not errores or not EXCEL_AVAILABLE:
        return None
    try:
        nombre = f"errores_{ips_nombre}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        path = dl_dir_empresa / nombre
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Errores"
        ws.append(["N° Factura", "Siniestro", "Error", "Fecha/Hora"])
        for e in errores:
            ws.append([e["factura"], e.get("siniestro_raw", ""), e.get("error", ""), e.get("timestamp", "")])
        wb.save(path)
        return nombre
    except Exception:
        return None


def crear_zip_final(dl_dir: Path, ips_nombre: str, zip_prefix: str, errores=None):
    """
    Empaqueta TODO lo de la carpeta de la IPS: Soportes/ (PDFs), Errores/
    (solo si hubo algún error), progreso.json y el reporte Excel.
    Si se pasan `errores`, también genera y embebe un Excel dedicado SOLO
    con las facturas fallidas (recargable directo al bot).
    """
    ips_dir = dl_dir / ips_nombre
    if not ips_dir.exists():
        return None
    zip_path = dl_dir / f"{zip_prefix}_{ips_nombre}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for f in ips_dir.rglob("*"):
            if f.is_file():
                zf.write(f, arcname=str(f.relative_to(ips_dir)))
        if errores:
            excel_solo_errores_nombre = generar_excel_solo_errores(ips_dir, ips_nombre, errores)
            if excel_solo_errores_nombre:
                excel_solo_errores_path = ips_dir / excel_solo_errores_nombre
                if excel_solo_errores_path.exists():
                    zf.write(excel_solo_errores_path, arcname=excel_solo_errores_nombre)
    # Ya terminó bien: cualquier ZIP parcial que haya quedado de intentos
    # anteriores queda obsoleto (el final ya lo incluye todo), se borra.
    for viejo in dl_dir.glob(f"{zip_prefix}_{ips_nombre}_PARCIAL_*.zip"):
        try:
            viejo.unlink()
        except Exception:
            pass
    return zip_path
    return zip_path


def generar_zip_parcial(job, empresa, dl_dir: Path, ips_dir: Path, ips_nombre: str, cfg, exitosas, errores):
    """Genera un ZIP con lo que se alcanzó a descargar hasta el momento (los
    PDFs ya en Soportes/, más un Excel parcial), para no perder el trabajo
    si el proceso se cae por un error antes de llegar al final normal."""
    if not ips_dir.exists():
        return
    if not exitosas and not errores:
        return  # nada que empacar todavía
    try:
        generar_reporte_excel(ips_dir, ips_nombre, exitosas, errores)
    except Exception as e:
        log(job, empresa, f"No se pudo generar el Excel parcial: {e}", "warn")
    try:
        # El nuevo parcial siempre es superconjunto del anterior (progreso.json
        # nunca repite lo ya descargado), así que el viejo queda redundante.
        for viejo in dl_dir.glob(f"{cfg['zip_prefix']}_{ips_nombre}_PARCIAL_*.zip"):
            try:
                viejo.unlink()
            except Exception:
                pass
        zip_path = dl_dir / f"{cfg['zip_prefix']}_{ips_nombre}_PARCIAL_{datetime.now().strftime('%Y%m%d_%H%M%S')}.zip"
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
            for f in ips_dir.rglob("*"):
                if f.is_file():
                    zf.write(f, arcname=str(f.relative_to(ips_dir)))
        log(job, empresa, f"📦 ZIP parcial generado con lo descargado hasta el momento: {zip_path.name}", "warn")
    except Exception as e:
        log(job, empresa, f"No se pudo generar el ZIP parcial: {e}", "error")


# ==================== LECTURA DEL EXCEL/CSV DE FACTURAS ====================
def leer_facturas_desde_archivo(filename: str, contenido: bytes):
    """Columna requerida: 'Siniestro' (17155/2026*3) — el portal nunca pide
    el N° de Factura para descargar, solo el siniestro. Si además viene una
    columna 'N° Factura' (o 'Factura') se usa para nombrar el PDF; si no,
    se deriva un identificador único a partir del propio siniestro."""
    filas, errores_formato = [], []
    nombre = filename.lower()

    def procesar_filas(iterador):
        filas_iter = list(iterador)
        idx_factura = None
        idx_siniestro = None
        start = 0
        for i, row in enumerate(filas_iter):
            if not row or all(c is None or c == "" for c in row):
                continue
            valores = [str(c).strip().lower() if c not in (None, "") else "" for c in row]
            if any("siniestro" in v for v in valores):
                for j, v in enumerate(valores):
                    if "factura" in v:
                        idx_factura = j
                    if "siniestro" in v:
                        idx_siniestro = j
                start = i + 1
            break
        if idx_siniestro is None:
            # Sin encabezado reconocible: se asume una sola columna con el siniestro
            idx_siniestro = 0

        for i, row in enumerate(filas_iter[start:], start=start + 1):
            if not row or all(c is None or c == "" for c in row):
                continue
            siniestro_raw = row[idx_siniestro] if len(row) > idx_siniestro else None
            if siniestro_raw in (None, ""):
                continue
            partes = parse_siniestro(siniestro_raw)
            if not partes:
                errores_formato.append({"fila": i, "factura": "", "siniestro": siniestro_raw})
                continue
            factura_raw = row[idx_factura] if idx_factura is not None and len(row) > idx_factura else None
            if factura_raw in (None, ""):
                factura = f"{partes['numero']}-{partes['anio']}-{partes['cuenta']}"
            else:
                factura = sanitizar_factura(factura_raw)
            filas.append({
                "factura": factura,
                "siniestro_raw": str(siniestro_raw).strip(),
                **partes,
            })

    if nombre.endswith(".csv"):
        for enc in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
            try:
                texto = contenido.decode(enc)
                break
            except (UnicodeDecodeError, LookupError):
                continue
        else:
            texto = contenido.decode("latin-1", errors="replace")
        reader = csv.reader(texto.splitlines())
        procesar_filas(reader)
    elif nombre.endswith((".xlsx", ".xls")):
        if not EXCEL_AVAILABLE:
            raise RuntimeError("openpyxl no instalado")
        wb = openpyxl.load_workbook(BytesIO(contenido), data_only=True)
        ws = wb.active
        procesar_filas(ws.iter_rows(values_only=True))
    else:
        raise ValueError("Formato no soportado. Usa CSV o Excel (.xlsx)")

    return filas, errores_formato


def generar_plantilla_ejemplo(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Facturas"
    ws.append(["Siniestro"])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.append(["17155/2026*3"])
    ws.column_dimensions["A"].width = 20
    wb.save(path)


# ==================== AUTOMATIZACIÓN (PLAYWRIGHT) ====================
# Mapa case-insensitive para XPath (incluye mayúsculas con tildes/ñ)
_CI_UPPER = "ÁÉÍÓÚÑABCDEFGHIJKLMNOPQRSTUVWXYZ"
_CI_LOWER = "áéíóúñabcdefghijklmnopqrstuvwxyz"


def _escribir_en_foco(page, value):
    """Teclea un valor en el campo que ya tiene foco, con eventos reales de teclado."""
    page.keyboard.press("Control+A")
    page.keyboard.press("Delete")
    page.keyboard.type(str(value), delay=40)
    page.keyboard.press("Tab")


def _escribir_en_locator(loc, value):
    """Escribe un valor en un locator ya encontrado (rápido, sin Tab final)."""
    try:
        loc.click()
    except Exception:
        pass
    try:
        loc.fill("")
    except Exception:
        pass
    try:
        loc.press_sequentially(str(value), delay=40)
    except Exception:
        pass


def _fallback_teclado(job, empresa, page, label_text, value, max_tabs=12):
    """
    Fallback rápido: si no encontramos el campo por locator, navegamos con
    Tab hasta encontrar un input cuyo label adyacente coincida con label_text.
    """
    log(job, empresa, f"  ⌨️  Fallback teclado para '{label_text}'...")
    needle = label_text.lower()
    try:
        page.locator("body").click(position={"x": 1, "y": 1})
    except Exception:
        pass
    page.wait_for_timeout(150)

    for i in range(max_tabs):
        try:
            info = page.evaluate(
                """() => {
                    const a = document.activeElement;
                    if (!a) return null;
                    let el = a, foundLabel = '';
                    for (let k = 0; k < 6 && el; k++) {
                        const lbl = el.querySelector && el.querySelector('label');
                        if (lbl && lbl.innerText) { foundLabel = lbl.innerText; break; }
                        el = el.parentElement;
                    }
                    if (!foundLabel && a.id) {
                        const lbl = document.querySelector(`label[for="${a.id}"]`);
                        if (lbl) foundLabel = lbl.innerText;
                    }
                    return {
                        tag: a.tagName,
                        placeholder: a.placeholder || '',
                        ariaLabel: a.getAttribute('aria-label') || '',
                        label: foundLabel,
                    };
                }"""
            )
        except Exception:
            info = None

        if info and info.get("tag") in ("INPUT", "TEXTAREA"):
            blob = " ".join([
                str(info.get("label") or ""),
                str(info.get("placeholder") or ""),
                str(info.get("ariaLabel") or ""),
            ]).lower()
            if needle in blob:
                _escribir_en_foco(page, value)
                log(job, empresa, f"  ⌨️  '{label_text}' alcanzado tras {i} Tab(s).")
                return True

        page.keyboard.press("Tab")
        page.wait_for_timeout(80)

    return False


def fill_by_label(job, empresa, page, label_text, value, timeout=8000):
    """
    Busca el <input> asociado a un texto de etiqueta y lo llena con tecleo real
    (press_sequentially en vez de .fill(), porque el portal valida con eventos
    de teclado reales). La contraseña se enmascara en el log.

    Performance: las estrategias se prueban en orden de probabilidad, con
    timeouts cortos. Caso común (XPath principal matchea) → ~1s.
    """
    valor_log = "•" * len(str(value)) if "contrase" in label_text.lower() else value
    log(job, empresa, f"  ✏️ Escribiendo en '{label_text}': {valor_log}")

    needle = label_text.lower()

    try:
        cand = page.get_by_label(label_text, exact=False).first
        cand.wait_for(state="visible", timeout=400)
        _escribir_en_locator(cand, value)
        return cand
    except Exception:
        pass

    xpath_principal = (
        f"xpath=//*[contains(translate(., '{_CI_UPPER}', '{_CI_LOWER}'), '{needle}')]"
        f"/following::input[1]"
    )
    try:
        cand = page.locator(xpath_principal).first
        cand.wait_for(state="visible", timeout=2500)
        _escribir_en_locator(cand, value)
        return cand
    except Exception:
        pass

    xpaths_alternativos = [
        f"xpath=//label[contains(translate(normalize-space(.), '{_CI_UPPER}', '{_CI_LOWER}'), '{needle}')]/following::input[1]",
        f"xpath=//input[@aria-label[contains(translate(., '{_CI_UPPER}', '{_CI_LOWER}'), '{needle}')]]",
        f"xpath=//*[contains(normalize-space(text()), '{label_text}')]/following::input[1]",
    ]
    for xp in xpaths_alternativos:
        try:
            cand = page.locator(xp).first
            cand.wait_for(state="visible", timeout=1200)
            _escribir_en_locator(cand, value)
            return cand
        except Exception:
            continue

    if _fallback_teclado(job, empresa, page, label_text, value, max_tabs=12):
        return None

    try:
        dbg_dir = BASE_DIR / "downloads" / empresa / "Errores"
        dbg_dir.mkdir(parents=True, exist_ok=True)
        safe = re.sub(r"[^A-Za-z0-9_]+", "_", label_text)[:40]
        shot = dbg_dir / f"debug_no_field_{safe}.png"
        html = dbg_dir / f"debug_no_field_{safe}.html"
        try:
            page.screenshot(path=str(shot), full_page=True)
        except Exception:
            pass
        try:
            html.write_text(page.content(), encoding="utf-8")
        except Exception:
            pass
        log(job, empresa,
            f"  ❌ No se encontró el campo '{label_text}'. Captura: {shot.name} | HTML: {html.name}",
            "error")
    except Exception:
        pass
    raise RuntimeError(f"No se encontró el campo '{label_text}' en el formulario.")


def run_automation(job: dict, empresa: str, usuario: str, password: str, ips_nombre: str, nit: str,
                    filas: list, download_path: str, ciudad_seleccionada: str = ""):
    from playwright.sync_api import sync_playwright

    cfg = EMPRESAS[empresa]
    job_state = job["state"]
    job_lock = job["lock"]

    dl_dir = Path(download_path)
    ips_dir = dl_dir / ips_nombre
    soportes_dir = ips_dir / "Soportes"
    soportes_dir.mkdir(parents=True, exist_ok=True)

    completadas = cargar_progreso(ips_dir)
    exitosas, errores = [], []
    errores_contados = set()
    zip_ya_generado = False
    with job_lock:
        job_state["stats"]["total"] = len(filas)

    ciudad = ciudad_seleccionada or cfg["ciudad_login"].get(nit, "VERIFICAR")
    if ciudad == "VERIFICAR":
        log(job, empresa, f"⚠️ No hay ciudad de login confirmada para el NIT {nit}. Revisa 'ciudad_login' en EMPRESAS.", "error")

    def procesar_lote(lote_filas, es_reintento=False):
        nonlocal exitosas, errores, errores_contados
        for fila in lote_filas:
            if job_state.get("stopping"):
                log(job, empresa, "🛑 Proceso detenido por el usuario.")
                return False
            factura = fila["factura"]
            if factura in completadas:
                continue

            # Cerrar cualquier modal residual antes de la siguiente factura
            try:
                ok_btn = page.locator('button:has-text("Ok")')
                if ok_btn.count() > 0 and ok_btn.is_visible():
                    ok_btn.click()
                    page.wait_for_selector('button:has-text("Ok")', state="hidden", timeout=3000)
            except Exception:
                pass
            page.wait_for_timeout(500)

            log(job, empresa, f"  → Procesando factura {factura}: {fila['numero']}/{fila['anio']}*{fila['cuenta']}"
                + (" (REINTENTO)" if es_reintento else ""))

            try:
                fill_by_label(job, empresa, page, "Número Siniestro", fila["numero"])
                fill_by_label(job, empresa, page, "Año Siniestro", fila["anio"])
                fill_by_label(job, empresa, page, "Número Cuenta", fila["cuenta"])

                destino = soportes_dir / f"{factura}.pdf"

                total_antes_resp = len(respuestas_detectadas)
                total_antes_dl = len(descargas_detectadas)

                # 1) Enviar -> spinner "Consultando..."
                page.click('button:has-text("Enviar")')

                # 2) Modal de éxito o de error
                loc_exito = page.get_by_text("exitosa")
                loc_error = page.get_by_text("Error", exact=True)
                loc_exito.or_(loc_error).first.wait_for(timeout=15000)

                if loc_error.count() > 0:
                    mensaje = ""
                    candidatos_mensaje = [
                        "xpath=//*[contains(normalize-space(text()),'Error')]/following::p[1]",
                        "xpath=//*[contains(normalize-space(text()),'Error')]/following-sibling::*[1]",
                        "xpath=//*[contains(normalize-space(text()),'Error')]/parent::*",
                    ]
                    for xp in candidatos_mensaje:
                        try:
                            texto = page.locator(xp).inner_text(timeout=1000).strip()
                            texto = texto.replace("Error", "").replace("Ok", "").strip()
                            if texto:
                                mensaje = texto
                                break
                        except Exception:
                            continue
                    log(job, empresa, f"  ⚠️ Modal de error del portal para factura {factura}: "
                        f"\"{mensaje or '(no se pudo leer el texto exacto, ver captura)'}\"", "warn")
                    try:
                        page.click('button:has-text("Ok")', timeout=2000)
                    except Exception:
                        pass
                    raise RuntimeError(f"El portal rechazó la solicitud: {mensaje or 'ver captura'}")

                # Éxito: 4 estrategias para capturar el PDF, en orden
                pdf_bytes = None
                nueva_pagina = None
                try:
                    with context.expect_page(timeout=4000) as page_info:
                        page.click('button:has-text("Ok")')
                    nueva_pagina = page_info.value
                except Exception:
                    nueva_pagina = None

                if nueva_pagina:
                    try:
                        for _ in range(20):
                            if nueva_pagina.url and nueva_pagina.url != "about:blank":
                                break
                            time.sleep(0.3)
                        pdf_url = nueva_pagina.url
                        try:
                            nueva_pagina.close()
                        except Exception:
                            pass
                        if pdf_url and pdf_url.startswith("http"):
                            resp_directa = context.request.get(pdf_url, timeout=60000)
                            if resp_directa.ok:
                                cuerpo = resp_directa.body()
                                if cuerpo[:4] == b"%PDF":
                                    pdf_bytes = cuerpo
                    except Exception:
                        pass

                page.wait_for_timeout(1200)

                if not pdf_bytes:
                    for resp in respuestas_detectadas[total_antes_resp:]:
                        try:
                            ct = (resp.headers.get("content-type") or "").lower()
                            if "pdf" in ct or "octet-stream" in ct:
                                cuerpo = resp.body()
                                if cuerpo[:4] == b"%PDF":
                                    pdf_bytes = cuerpo
                                    break
                        except Exception:
                            continue

                if not pdf_bytes:
                    for resp in respuestas_detectadas[total_antes_resp:]:
                        try:
                            ct = (resp.headers.get("content-type") or "").lower()
                            if "json" in ct or "text" in ct or ct == "":
                                texto = resp.text()
                                for m in re.finditer(r'"([A-Za-z0-9+/]{500,}={0,2})"', texto):
                                    try:
                                        candidato = base64.b64decode(m.group(1))
                                    except Exception:
                                        continue
                                    if candidato[:4] == b"%PDF":
                                        pdf_bytes = candidato
                                        break
                                if pdf_bytes:
                                    break
                        except Exception:
                            continue

                if pdf_bytes:
                    destino.write_bytes(pdf_bytes)
                elif len(descargas_detectadas) > total_antes_dl:
                    descargas_detectadas[-1].save_as(str(destino))
                else:
                    (ips_dir / "Errores").mkdir(parents=True, exist_ok=True)
                    debug_path = ips_dir / "Errores" / f"debug_red_{factura}.txt"
                    try:
                        lineas = []
                        for resp in respuestas_detectadas[total_antes_resp:]:
                            try:
                                ct = resp.headers.get("content-type") or ""
                                linea = f"{resp.status} | {ct} | {resp.url}"
                                if "json" in ct.lower() or "text" in ct.lower():
                                    try:
                                        preview = resp.text()[:300].replace("\n", " ")
                                        linea += f"\n    preview: {preview}"
                                    except Exception:
                                        pass
                                lineas.append(linea)
                            except Exception:
                                pass
                        debug_path.write_text("\n".join(lineas), encoding="utf-8")
                    except Exception:
                        pass
                    raise RuntimeError(
                        f"No se pudo capturar el PDF tras el clic en 'Ok'. Diagnóstico guardado en {debug_path.name}."
                    )

                exitosas.append({
                    "factura": factura, "numero": fila["numero"], "anio": fila["anio"],
                    "cuenta": fila["cuenta"], "archivo": destino.name,
                    "timestamp": datetime.now(timezone.utc).isoformat(),
                })
                completadas.add(factura)
                guardar_progreso(ips_dir, exitosas, meta={
                    "identidad": job["state"].get("identidad"),
                    "ips": job["state"].get("ips_identity"),
                    "ejecucion_id": job["state"].get("ejecucion_id"),
                    "ips_nombre": ips_nombre,
                    "aseguradora": EMPRESAS[empresa]["nombre"],
                    "sede": ciudad,
                })
                with job_lock:
                    job_state["stats"]["descargadas"] += 1
                    job_state["descargas_exitosas"] = exitosas
                    if factura in errores_contados:
                        errores_contados.discard(factura)
                        job_state["stats"]["errores"] = max(0, job_state["stats"]["errores"] - 1)
                log(job, empresa, f"  ✅ Factura {factura} descargada.")

            except Exception as e:
                (ips_dir / "Errores").mkdir(parents=True, exist_ok=True)
                captura = ips_dir / "Errores" / f"error_{factura}.png"
                try:
                    page.screenshot(path=str(captura))
                except Exception:
                    captura = None
                errores.append({
                    "factura": factura, "siniestro_raw": fila["siniestro_raw"],
                    "error": str(e), "captura": captura.name if captura else "",
                    "timestamp": datetime.now(timezone.utc).isoformat(),
                })
                if factura not in errores_contados:
                    errores_contados.add(factura)
                    with job_lock:
                        job_state["stats"]["errores"] += 1
                        job_state["errores_detalle"] = errores
                log(job, empresa, f"  ❌ Error en factura {factura}: {e}", "error")
                try:
                    historial_db.registrar_error_ejecucion(
                        job_state.get("ejecucion_id"), factura, empresa, EMPRESAS.get(empresa, {}).get("nombre"),
                        job_state.get("ips_identity"), "descarga_persistente", str(e), "descarga_factura",
                        True, 1, "fallida",
                    )
                except Exception as e2:
                    log(job, empresa, f"⚠️ No se pudo registrar el error de la factura {factura} en el historial: {e2}", "warn")

                try:
                    ok_btn = page.locator('button:has-text("Ok")')
                    if ok_btn.count() > 0 and ok_btn.is_visible():
                        ok_btn.click()
                        page.wait_for_selector('button:has-text("Ok")', state="hidden", timeout=3000)
                except Exception:
                    pass
                page.wait_for_timeout(1000)
        return True

    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True, args=["--no-sandbox", "--disable-dev-shm-usage"])  # Railway: sin pantalla
            context = browser.new_context(accept_downloads=True, viewport={"width": 1500, "height": 900})
            page = context.new_page()
            job["browser"] = browser

            descargas_detectadas = []
            page.on("download", lambda d: descargas_detectadas.append(d))
            respuestas_detectadas = []
            page.on("response", lambda r: respuestas_detectadas.append(r))

            log(job, empresa, f"🔐 Iniciando sesión — Usuario: {usuario} / Ciudad: {ciudad}")
            page.goto(cfg["login_url"], wait_until="domcontentloaded", timeout=60000)
            try:
                page.wait_for_load_state("load", timeout=30000)
            except Exception:
                pass
            page.wait_for_timeout(1500)

            fill_by_label(job, empresa, page, "NIT", usuario)

            estrategias_ciudad = [
                "xpath=//label[contains(translate(normalize-space(.), 'ÁÉÍÓÚÑABCDEFGHIJKLMNOPQRSTUVWXYZ', 'áéíóúñabcdefghijklmnopqrstuvwxyz'), 'ciudad')]/following::input[1]",
                "xpath=//*[contains(translate(., 'ÁÉÍÓÚÑABCDEFGHIJKLMNOPQRSTUVWXYZ', 'áéíóúñabcdefghijklmnopqrstuvwxyz'), 'ciudad')]/following::input[1]",
                "xpath=//*[contains(normalize-space(text()), 'Ciudad')]/following::input[1]",
            ]
            loc_ciudad = None
            for xp in estrategias_ciudad:
                try:
                    cand = page.locator(xp).first
                    cand.wait_for(state="visible", timeout=8000)
                    loc_ciudad = cand
                    break
                except Exception:
                    continue

            if loc_ciudad is None:
                raise RuntimeError("No se encontró el campo 'Ciudad' en el formulario.")

            loc_ciudad.click()
            termino_busqueda = ciudad.split(" - ")[0]
            loc_ciudad.press_sequentially(termino_busqueda, delay=60)
            try:
                page.wait_for_selector(f"text={ciudad}", timeout=8000)
                page.click(f"text={ciudad}")
            except Exception:
                log(job, empresa, f"⚠️ No apareció la sugerencia de ciudad '{ciudad}'. Verifica el texto exacto.", "warn")

            fill_by_label(job, empresa, page, "Contrase", password)
            page.click('button:has-text("Ingresar")')

            url_post_login = None
            patrones_post_login = ("**/dashboard/**", "**/protected/**", "**/inicio/**",
                                   "**/home/**", "**/main/**", "**/urls", "**/urls/**")
            for patron in patrones_post_login:
                try:
                    page.wait_for_url(patron, timeout=6000)
                    url_post_login = page.url
                    break
                except Exception:
                    continue
            if not url_post_login:
                log(job, empresa,
                    f"  ⚠️ URL post-login no reconocida ({page.url}); continúo de todos modos.",
                    "warn")
                page.wait_for_timeout(4000)
            else:
                log(job, empresa, f"✅ Sesión iniciada correctamente. URL: {url_post_login}")

            try:
                page.wait_for_selector("text=BIENVENIDO", timeout=15000)
            except Exception:
                page.wait_for_timeout(3000)

            def _navegar_por_teclado():
                for _ in range(11):
                    page.keyboard.press("Tab")
                    page.wait_for_timeout(120)
                page.keyboard.press("Space")
                page.wait_for_timeout(300)
                for _ in range(2):
                    page.keyboard.press("Tab")
                    page.wait_for_timeout(120)
                page.keyboard.press("Enter")
                page.wait_for_selector("text=GENERAR LIQUIDACION SOAT", timeout=8000)

            menu_desplegado = False
            try:
                if page.get_by_text("GENERAR LIQUIDACION SOAT").first.is_visible(timeout=2000):
                    menu_desplegado = True
                    log(job, empresa, "📄 Módulo de liquidación ya estaba visible.")
            except Exception:
                pass

            if not menu_desplegado:
                try:
                    _navegar_por_teclado()
                    menu_desplegado = True
                    log(job, empresa, "📄 Módulo de liquidación abierto (navegación por teclado).")
                except Exception:
                    log(job, empresa,
                        "  → La navegación por teclado no llegó al formulario, probando por clic en el menú lateral...",
                        "warn")

            if not menu_desplegado:
                candidatos_menu = page.get_by_text("Generación De Soportes De Notificación", exact=True)
                total_candidatos = candidatos_menu.count()
                log(job, empresa,
                    f"  🔎 Encontré {total_candidatos} coincidencia(s) para 'Generación De Soportes De Notificación'.")

                for i in range(total_candidatos):
                    for objetivo in (
                        candidatos_menu.nth(i),
                        candidatos_menu.nth(i).locator("xpath=ancestor::li[1]"),
                        candidatos_menu.nth(i).locator("xpath=ancestor::a[1]"),
                        candidatos_menu.nth(i).locator("xpath=ancestor::button[1]"),
                        candidatos_menu.nth(i).locator(
                            "xpath=ancestor::*[contains(@class,'menu') or contains(@class,'nav') or contains(@class,'item')][1]"
                        ),
                    ):
                        try:
                            objetivo.scroll_into_view_if_needed(timeout=2000)
                        except Exception:
                            pass
                        try:
                            objetivo.click(timeout=4000, force=True)
                            try:
                                page.wait_for_selector("text=Soporte De Liquidación",
                                                        state="visible", timeout=5000)
                                menu_desplegado = True
                                log(job, empresa,
                                    f"  ✓ Submenú 'Soporte De Liquidación' desplegado tras clic en coincidencia #{i+1}.")
                                break
                            except Exception:
                                try:
                                    if page.get_by_text("GENERAR LIQUIDACION SOAT").first.is_visible(timeout=1500):
                                        menu_desplegado = True
                                        log(job, empresa,
                                            f"  ✓ Módulo 'GENERAR LIQUIDACION SOAT' visible tras clic en coincidencia #{i+1}.")
                                        break
                                except Exception:
                                    pass
                        except Exception:
                            continue
                    if menu_desplegado:
                        break

                if menu_desplegado:
                    try:
                        if page.get_by_text("Soporte De Liquidación").first.is_visible(timeout=2000):
                            page.get_by_text("Soporte De Liquidación").first.click(timeout=10000)
                    except Exception:
                        log(job, empresa,
                            "  ⚠️ No encontré 'Soporte De Liquidación' como sub-item clickeable; "
                            "puede que el item padre ya haya abierto el módulo directo.",
                            "warn")
                    try:
                        page.wait_for_selector("text=GENERAR LIQUIDACION SOAT", timeout=10000)
                        log(job, empresa, "📄 Módulo de liquidación abierto (clic por menú lateral).")
                    except Exception:
                        try:
                            if page.get_by_text("GENERAR LIQUIDACION SOAT").first.is_visible(timeout=2000):
                                log(job, empresa, "📄 Módulo de liquidación ya visible.")
                            else:
                                raise RuntimeError("No se confirmó 'GENERAR LIQUIDACION SOAT' tras la navegación.")
                        except Exception:
                            raise

            if not menu_desplegado:
                (ips_dir / "Errores").mkdir(parents=True, exist_ok=True)
                debug_path = ips_dir / "Errores" / "debug_menu.html"
                try:
                    debug_path.write_text(page.content(), encoding="utf-8")
                    log(job, empresa, f"  → Se guardó el HTML de la página en {debug_path} para diagnóstico.", "warn")
                except Exception:
                    pass
                raise RuntimeError("No se pudo llegar al módulo de Liquidación SOAT (ni por teclado ni por menú).")

            log(job, empresa, f"📦 Iniciando primer ciclo de descarga con {len(filas)} facturas.")
            procesar_lote(filas, es_reintento=False)

            fallidas = [f for f in filas if f["factura"] not in completadas]
            if fallidas and not job_state.get("stopping"):
                log(job, empresa, f"🔄 Reintentando {len(fallidas)} facturas que fallaron en el primer ciclo.")
                procesar_lote(fallidas, es_reintento=True)

            generar_reporte_excel(ips_dir, ips_nombre, exitosas, errores)
            zip_final_path = crear_zip_final(dl_dir, ips_nombre, cfg["zip_prefix"], errores=errores)
            zip_ya_generado = True
            if zip_final_path:
                try:
                    with job_lock:
                        job_state["zip_url"] = f"/downloads/{empresa}/{zip_final_path.name}"
                except Exception:
                    pass

            # Excel aparte, solo de errores, para descarga automática desde
            # el frontend — nunca debe poder tumbar el proceso ya exitoso.
            try:
                nombre_excel_errores = generar_excel_solo_errores(dl_dir / empresa, ips_nombre, errores)
                with job_lock:
                    job_state["errores_excel_url"] = (
                        f"/downloads/{empresa}/{nombre_excel_errores}" if nombre_excel_errores else None
                    )
            except Exception:
                pass

            try:
                browser.close()
            except Exception:
                pass

    except Exception as e:
        log(job, empresa, f"❌ Error general: {e}", "error")
        with job_lock:
            job_state["error"] = str(e)
    finally:
        # Red de seguridad final: sin importar CÓMO se salió de la función
        # (éxito, error, o Detener manual del usuario), si todavía no se
        # generó ningún ZIP, se genera uno parcial aquí con lo que se haya
        # descargado hasta el momento.
        if not zip_ya_generado:
            generar_zip_parcial(job, empresa, dl_dir, ips_dir, ips_nombre, cfg, exitosas, errores)
        historial_db.cerrar_ejecucion(
            job_state.get("ejecucion_id"), "error" if job_state.get("error") else ("cancelada" if job_state.get("stopping") else "completada"),
            total_detectadas=len(filas), total_procesadas=len(exitosas) + len(errores),
            total_exitosas=len(exitosas), total_fallidas=len(errores), total_redescargadas=0,
        )
        job["browser"] = None
        with job_lock:
            job_state["running"] = False
            job_state["finished"] = True
            job_state["stopping"] = False
        concurrency.registrar_fin(empresa)


# ==================== RUTAS FLASK ====================
@bp.route("/<any(estado,sura):empresa>")
def index(empresa):
    cfg = empresa_o_404(empresa)
    return render_template("estado_sura_index.html", empresa=empresa, cfg=cfg,
                            otras=[e for e in EMPRESAS if e != empresa])


@bp.route("/plantilla")
def descargar_plantilla():
    path = BASE_DIR / "plantilla_facturas.xlsx"
    generar_plantilla_ejemplo(path)
    return send_file(path, as_attachment=True, download_name="plantilla_facturas.xlsx")


@bp.route("/api/<empresa>/upload", methods=["POST"])
def upload_facturas(empresa):
    """Validación/preview sin estado: solo confirma que el archivo se puede
    leer y cuántas facturas trae. El archivo se vuelve a enviar (y a
    procesar de verdad) junto con /start, para que quede atado a la cuenta
    correcta (empresa_id) de forma atómica."""
    empresa_o_404(empresa)
    if "file" not in request.files:
        return jsonify({"ok": False, "error": "No se envió ningún archivo"}), 400
    file = request.files["file"]
    if file.filename == "":
        return jsonify({"ok": False, "error": "Archivo vacío"}), 400
    try:
        filas, errores_formato = leer_facturas_desde_archivo(file.filename, file.read())
        if not filas:
            return jsonify({"ok": False, "error": "No se encontraron filas válidas (revisa el formato del siniestro)"}), 400
        return jsonify({"ok": True, "count": len(filas), "omitidas": len(errores_formato)})
    except Exception as e:
        return jsonify({"ok": False, "error": f"Error al procesar archivo: {e}"}), 500



@bp.route("/api/<empresa>/start", methods=["POST"])
def start_job(empresa):
    cfg = empresa_o_404(empresa)
    usuario = request.form.get("usuario", "").strip()
    password = request.form.get("password", "").strip()
    custom_path = request.form.get("download_path", "").strip()
    archivo = request.files.get("file")
    manual_json = request.form.get("manual_entries", "").strip()
    identidad = validar_identidad(request.form.get("identidad"))
    ciudad_seleccionada = request.form.get("ciudad_login", "").strip()

    if not identidad:
        return jsonify({"ok": False, "error": "Selecciona quién eres antes de iniciar el proceso.", "campo": "identidad"}), 400

    if not usuario or not password:
        return jsonify({"ok": False, "error": "Faltan usuario y/o contraseña"}), 400

    _, nit_usuario_sede = extraer_ips_desde_usuario(usuario, cfg["mapa_ips"])
    sedes_disponibles = cfg.get("sedes_login", {}).get(nit_usuario_sede, [])
    if sedes_disponibles and ciudad_seleccionada not in sedes_disponibles:
        return jsonify({"ok": False, "error": "Selecciona una sede válida antes de iniciar el proceso.", "campo": "ciudad_login"}), 400

    manual_filas = []
    errores_formato = []
    if manual_json:
        try:
            manual_data = json.loads(manual_json)
        except Exception:
            return jsonify({"ok": False, "error": "Formato inválido en las facturas agregadas manualmente"}), 400
        for item in manual_data:
            # Acepta tanto texto plano (solo el siniestro) como el formato
            # viejo {"factura":..., "siniestro":...} por compatibilidad.
            if isinstance(item, dict):
                siniestro_raw = str(item.get("siniestro", "")).strip()
                factura_raw = str(item.get("factura", "")).strip()
            else:
                siniestro_raw = str(item).strip()
                factura_raw = ""
            if not siniestro_raw:
                continue
            partes = parse_siniestro(siniestro_raw)
            if not partes:
                errores_formato.append({"fila": "manual", "factura": factura_raw, "siniestro": siniestro_raw})
                continue
            factura = sanitizar_factura(factura_raw) if factura_raw else f"{partes['numero']}-{partes['anio']}-{partes['cuenta']}"
            manual_filas.append({
                "factura": factura,
                "siniestro_raw": siniestro_raw,
                **partes,
            })

    if not archivo or archivo.filename == "":
        if not manual_filas:
            return jsonify({"ok": False, "error": "Falta el Excel de facturas, o agrega al menos una manualmente"}), 400
        filas = []
    else:
        try:
            filas, errores_excel = leer_facturas_desde_archivo(archivo.filename, archivo.read())
        except Exception as e:
            return jsonify({"ok": False, "error": f"Error al procesar archivo: {e}"}), 500
        errores_formato.extend(errores_excel)

    filas = filas + manual_filas

    empresa_id = resolve_empresa_id(usuario)
    job = get_or_create_job(empresa, empresa_id)

    with job["lock"]:
        if job["state"]["running"]:
            return jsonify({"ok": False, "error": "Ya hay un proceso en ejecución para esta cuenta"}), 409
    if count_running_jobs(empresa) >= MAX_CONCURRENT_JOBS:
        return jsonify({"ok": False, "error": f"Se alcanzó el máximo de {MAX_CONCURRENT_JOBS} procesos simultáneos para esta empresa. Intenta de nuevo en unos minutos."}), 429
    if not concurrency.puede_iniciar():
        return jsonify({"ok": False, "error": "El panel está al máximo de descargas simultáneas entre todas las aseguradoras en este momento. Intenta de nuevo en unos minutos."}), 429

    if not filas:
        return jsonify({"ok": False, "error": "No se encontraron filas válidas (revisa el formato del siniestro)"}), 400

    ips_nombre, nit = extraer_ips_desde_usuario(usuario, cfg["mapa_ips"])
    ips_nit_manual = normalizar_nit(request.form.get("ips_nit_manual", ""))
    if ips_nit_manual:
        nit = ips_nit_manual
        ips_nombre = resolver(nit=ips_nit_manual).get("nombre_estandar", ips_nombre)
    elif resolver(nit=nit, nombre_detectado=ips_nombre).get("metodo") == "NO_IDENTIFICADA":
        return jsonify({"ok": False, "requiere_seleccion_ips": True,
                         "catalogo_ips": listar_catalogo_por_responsable()})

    confirmar_duplicados = str(request.form.get("confirmar_duplicados", "")).lower() in ("true", "1", "on", "si", "sí")
    decision_redescarga = request.form.get("decision_redescarga", "")
    try:
        facturas_redescarga = {str(v) for v in json.loads(request.form.get("facturas_redescarga", "[]"))}
    except Exception:
        facturas_redescarga = set()
    try:
        if nit:
            ya_descargadas = historial_db.buscar_ya_descargadas_por_nit(
                cfg["nombre"], nit, [f["factura"] for f in filas]
            )
        else:
            ya_descargadas = historial_db.buscar_ya_descargadas(
                cfg["nombre"], ips_nombre, [f["factura"] for f in filas]
            )
    except Exception as e:
        log(job, empresa, f"⚠️ No se pudo consultar el historial de duplicados: {e}", "warn")
        ya_descargadas = {}
    if ya_descargadas:
        if not confirmar_duplicados:
            return jsonify({
                "ok": False, "requiere_confirmacion": True,
                "ya_descargadas": ya_descargadas, "total_filas": len(filas),
            })
        if decision_redescarga == "ninguna":
            filas = [f for f in filas if f["factura"] not in ya_descargadas]
        elif decision_redescarga == "seleccionadas":
            filas = [f for f in filas if f["factura"] not in ya_descargadas or f["factura"] in facturas_redescarga]

    with job["lock"]:
        job["state"].update({
            "running": True, "finished": False, "error": None,
            "stats": {"total": 0, "descargadas": 0, "errores": 0},
            "errores_detalle": [], "descargas_exitosas": [],
            "errores_excel_url": None, "zip_url": None,
        })
    job["ips_nombre"] = ips_nombre
    dl_dir_empresa = DOWNLOAD_DIR / empresa
    dl_path = custom_path if custom_path else str(dl_dir_empresa)
    if custom_path:
        registro_rutas.registrar_ruta(EMPRESAS[empresa]["nombre"], custom_path)

    for ef in errores_formato:
        log(job, empresa, f"⚠️ Fila {ef['fila']}: siniestro '{ef['siniestro']}' con formato inválido, se omite.", "warn")
    log(job, empresa, f"🚀 Proceso iniciado | IPS detectada: {ips_nombre} (NIT {nit}) | {len(filas)} facturas")

    job["state"]["identidad"] = identidad
    job["state"]["ips_identity"] = resolver(nit=nit, nombre_detectado=ips_nombre)
    job["state"]["ips_nit"] = nit
    job["state"]["ejecucion_id"] = historial_db.iniciar_ejecucion(
        identidad, empresa, cfg["nombre"], job["state"]["ips_identity"], None, dl_path
    )

    concurrency.registrar_inicio(empresa)
    t = threading.Thread(target=run_automation,
                          args=(job, empresa, usuario, password, ips_nombre, nit, filas, dl_path,
                                ciudad_seleccionada),
                          daemon=True)
    t.start()
    return jsonify({"ok": True, "ips": ips_nombre, "nit": nit, "empresa_id": empresa_id,
                     "download_path": dl_path, "total": len(filas)})


@bp.route("/api/<empresa>/stop", methods=["POST"])
def stop_job_route(empresa):
    empresa_o_404(empresa)
    empresa_id = request.args.get("empresa_id", "")
    job = get_job_or_none(empresa, empresa_id)
    if not job or not job["state"]["running"]:
        return jsonify({"ok": False, "message": "No hay proceso en ejecución"}), 400
    stop_job(job, empresa)
    return jsonify({"ok": True, "message": "Deteniendo proceso..."})


@bp.route("/api/<empresa>/reset", methods=["POST"])
def reset_job_route(empresa):
    empresa_o_404(empresa)
    data = request.get_json(silent=True) or request.form or {}
    ips = data.get("ips", "").strip()
    empresa_id = data.get("empresa_id", "").strip()
    job = get_job_or_none(empresa, empresa_id)

    if job and job["state"]["running"]:
        stop_job(job, empresa)
        # Espera activa (máx ~10s) hasta que el hilo confirme running=False
        # antes de borrar progreso y resetear estado -- antes solo se
        # esperaba un sleep(2) fijo, que podía no ser suficiente y dejar
        # el navegador de Playwright todavía cerrándose.
        for _ in range(20):
            time.sleep(0.5)
            with job["lock"]:
                if not job["state"]["running"]:
                    break

    if ips:
        progreso_file = DOWNLOAD_DIR / empresa / ips / "progreso.json"
        if progreso_file.exists():
            try:
                n_migradas = historial_db.migrar_progreso_json(progreso_file.parent)
                progreso_file.unlink()
                if job:
                    log(job, empresa, f"🗑️ Progreso eliminado: {progreso_file} ({n_migradas} factura(s) archivadas en el historial)")
            except Exception as e:
                if job:
                    log(job, empresa, f"⚠️ Error al borrar {progreso_file}: {e}", "warn")

    if job:
        reset_state(job)
    return jsonify({"ok": True, "message": "Estado reiniciado y progreso eliminado."})


@bp.route("/api/<empresa>/status")
def get_status(empresa):
    empresa_o_404(empresa)
    empresa_id = request.args.get("empresa_id", "")
    if empresa_id:
        job = get_job_or_none(empresa, empresa_id)
        if not job:
            return jsonify({"running": False, "finished": False, "error": None,
                             "stats": {"total": 0, "descargadas": 0, "errores": 0}, "logs": []})
        with job["lock"]:
            js = job["state"]
            return jsonify({
                "running": js["running"], "finished": js["finished"], "error": js["error"],
                "stats": js["stats"], "logs": js["logs"][-200:],
                "errores_excel_url": js.get("errores_excel_url"),
                "zip_url": js.get("zip_url"),
            })
    # Sin empresa_id (ej: el panel principal): vista agregada de todas las
    # cuentas activas para esta empresa, útil para un vistazo rápido.
    with jobs_registry_lock:
        activos = [j for k, j in jobs.items() if k.startswith(f"{empresa}:") and j["state"]["running"]]
        stats = {
            "total": sum(j["state"]["stats"]["total"] for j in activos),
            "descargadas": sum(j["state"]["stats"]["descargadas"] for j in activos),
            "errores": sum(j["state"]["stats"]["errores"] for j in activos),
        }
        procesos = [{
            "ips": j.get("ips_nombre") or "Detectando IPS...",
            "stats": j["state"]["stats"],
        } for j in activos]
        return jsonify({"running": len(activos) > 0, "finished": False, "error": None,
                         "stats": stats, "logs": [], "procesos": procesos})


@bp.route("/api/<empresa>/logs")
def get_logs(empresa):
    empresa_o_404(empresa)
    since = int(request.args.get("since", 0))
    empresa_id = request.args.get("empresa_id", "")
    job = get_job_or_none(empresa, empresa_id)
    if not job:
        return jsonify({"logs": []})
    with job["lock"]:
        return jsonify({"logs": job["state"]["logs"][since:]})


@bp.route("/api/<empresa>/logs", methods=["DELETE"])
def clear_logs(empresa):
    empresa_o_404(empresa)
    empresa_id = request.args.get("empresa_id", "")
    job = get_job_or_none(empresa, empresa_id)
    if job:
        with job["lock"]:
            job["state"]["logs"] = []
    return jsonify({"ok": True})


@bp.route("/api/<empresa>/clear", methods=["POST"])
def clear_panel(empresa):
    """
    Resetea TODO el panel de esta cuenta (log, stats, badge) a su estado
    inicial. NO toca progreso.json en disco — eso lo hace el botón
    'Reiniciar'/'Limpiar progreso', que sí borra el historial de
    reanudación. Solo afecta a la cuenta indicada (empresa+NIT), nunca a
    otra cuenta ni a la otra empresa.
    """
    empresa_o_404(empresa)
    empresa_id = request.args.get("empresa_id", "") or (request.get_json(silent=True) or {}).get("empresa_id", "")
    job = get_job_or_none(empresa, empresa_id)
    if not job:
        return jsonify({"ok": True})
    with job["lock"]:
        if job["state"]["running"]:
            return jsonify({"ok": False, "error": "No se puede limpiar mientras hay un proceso en ejecución. Detén primero."}), 409
        job["state"].update(_estado_inicial())
    return jsonify({"ok": True})


@bp.route("/api/<empresa>/files")
def list_files(empresa):
    cfg = empresa_o_404(empresa)
    ips = request.args.get("ips", "")
    dl_dir_empresa = DOWNLOAD_DIR / empresa
    files = []
    if dl_dir_empresa.exists():
        for f in sorted(dl_dir_empresa.glob("*.zip")):
            if not ips or f.name.startswith(f"{cfg['zip_prefix']}_{ips}_"):
                files.append({"name": f.name, "size": f.stat().st_size})
    return jsonify({"files": files})


@bp.route("/api/<empresa>/files", methods=["DELETE"])
def delete_all_files(empresa):
    cfg = empresa_o_404(empresa)
    ips = request.args.get("ips", "")
    dl_dir_empresa = DOWNLOAD_DIR / empresa
    try:
        eliminados = 0
        if dl_dir_empresa.exists():
            for f in list(dl_dir_empresa.glob("*.zip")):
                if not ips or f.name.startswith(f"{cfg['zip_prefix']}_{ips}_"):
                    f.unlink()
                    eliminados += 1
        log(None, empresa, f"🗑️ ZIPs eliminados: {eliminados}")
        return jsonify({"ok": True, "message": f"Se eliminaron {eliminados} ZIP(s).", "eliminados": eliminados})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500


@bp.route("/downloads/<empresa>/<path:filename>")
def download_file(empresa, filename):
    empresa_o_404(empresa)
    return send_from_directory(DOWNLOAD_DIR / empresa, filename, as_attachment=True)


@bp.route("/api/<empresa>/progreso")
def get_progreso(empresa):
    empresa_o_404(empresa)
    ips = request.args.get("ips", "")
    if not ips:
        return jsonify({"ok": False, "error": "Se requiere el parámetro 'ips'"}), 400
    progreso_path = DOWNLOAD_DIR / empresa / ips / "progreso.json"
    if not progreso_path.exists():
        return jsonify({"ok": True, "completadas": [], "ips": ips, "mensaje": "Aún no hay facturas completadas"})
    try:
        data = json.load(open(progreso_path, encoding="utf-8"))
        completadas = data.get("completadas", [])
        return jsonify({"ok": True, "completadas": completadas, "cantidad": len(completadas),
                         "ips": ips, "actualizado": data.get("actualizado", "")})
    except Exception as e:
        return jsonify({"ok": False, "error": f"Error al leer progreso: {e}"}), 500


@bp.route("/api/<empresa>/exportar_progreso")
def exportar_progreso_excel(empresa):
    empresa_o_404(empresa)
    ips = request.args.get("ips", "")
    if not ips:
        return jsonify({"ok": False, "error": "Se requiere el parámetro 'ips'"}), 400
    if not EXCEL_AVAILABLE:
        return jsonify({"ok": False, "error": "openpyxl no instalado"}), 500
    progreso_path = DOWNLOAD_DIR / empresa / ips / "progreso.json"
    if not progreso_path.exists():
        return jsonify({"ok": False, "error": f"No se encontró progreso.json para '{ips}'"}), 404
    try:
        data = json.load(open(progreso_path, encoding="utf-8"))
        completadas = data.get("completadas", [])
        actualizado = data.get("actualizado", "")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Facturas completadas"
        ws.append(["N° Factura", "Fecha de completado"])
        for factura in completadas:
            ws.append([factura, actualizado])
        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 30
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return send_file(output, as_attachment=True,
                          download_name=f"progreso_{empresa}_{ips}.xlsx",
                          mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except Exception as e:
        return jsonify({"ok": False, "error": f"Error al generar Excel: {e}"}), 500
